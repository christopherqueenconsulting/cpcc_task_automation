#  Copyright (c) 2026. Christopher Queen Consulting LLC (http://www.ChristopherQueenConsulting.com/)
"""Score-matched closings, the no-code floor for "Grade anyway", and the result card."""

import json

import pytest

from cqc_cpcc.rubric_models import CriterionResult, DetectedError, RubricAssessmentResult, RequirementResult
from cqc_cpcc.student_feedback_builder import (
    CLOSING_BANDS,
    CLOSING_NO_WORK,
    build_student_feedback,
    closing_for_result,
)


def _result(earned, possible=30, **extra):
    return RubricAssessmentResult(
        rubric_id="csc134_cpp_exam_rubric", rubric_version="3.1",
        total_points_possible=possible, total_points_earned=earned,
        criteria_results=[CriterionResult(criterion_id="program_performance", criterion_name="Program Performance",
                                          points_possible=possible, points_earned=earned,
                                          selected_level_label="Average", feedback="Works | mostly.")],
        overall_feedback="Instructor note.",
        detected_errors=[DetectedError(code="X", name="Off by one", severity="minor", description="Loop bound.")],
        requirement_results=[RequirementResult(requirement_id="R1", status="met", evidence="main reads input")],
        **extra,
    )


@pytest.mark.unit
@pytest.mark.parametrize("earned, expected_floor", [(30, 90), (27, 90), (25, 80), (22, 70), (19, 60), (4.5, 0), (0, 0)])
def test_closing_matches_the_score_band(earned, expected_floor):
    expected = dict(CLOSING_BANDS)[expected_floor]
    assert closing_for_result(_result(earned)) == expected


@pytest.mark.unit
def test_no_work_closing_and_no_old_closing():
    flagged = _result(0, needs_review=True, validity_status="wrong_type", validity_reason="docx")
    text = build_student_feedback(flagged)
    assert text.rstrip().endswith(CLOSING_NO_WORK)
    assert "Keep up the good work" not in build_student_feedback(_result(5))


@pytest.mark.unit
def test_closings_never_promise_a_next_assignment_or_resubmission():
    for text in [t for _, t in CLOSING_BANDS] + [CLOSING_NO_WORK]:
        lowered = text.lower()
        assert "next assignment" not in lowered and "next project" not in lowered
        assert "resubmit" not in lowered and "re-submit" not in lowered


@pytest.mark.unit
def test_greeting_without_a_name_is_plain_hello():
    assert build_student_feedback(_result(20), student_name=None).startswith("Hello,")


@pytest.mark.unit
def test_no_code_floor_zeroes_a_docx_lab_report_but_keeps_feedback(tmp_path):
    from cqc_cpcc.rubric_config import get_rubric_by_id
    from cqc_cpcc.rubric_grading import apply_no_code_floor

    rubric = get_rubric_by_id("csc134_cpp_exam_rubric")
    report = tmp_path / "lab 4 super positon.txt"
    report.write_text("Lab report: superposition of waves. We measured amplitudes.")
    graded = _result(4.5)
    floored = apply_no_code_floor(graded, rubric, {report.name: str(report)})
    assert floored.total_points_earned == 0
    assert floored.validity_status == "wrong_type" and floored.needs_review
    assert floored.overall_feedback == "Instructor note."
    assert any(e.name == "Off by one" for e in floored.detected_errors)
    assert floored.requirement_results == graded.requirement_results


@pytest.mark.unit
def test_no_code_floor_keeps_real_code_found_in_a_text_file(tmp_path):
    from cqc_cpcc.rubric_config import get_rubric_by_id
    from cqc_cpcc.rubric_grading import apply_no_code_floor

    rubric = get_rubric_by_id("csc134_cpp_exam_rubric")
    src = tmp_path / "answer.txt"
    src.write_text("#include <iostream>\nint main() { std::cout << 1; return 0; }\n")
    graded = _result(18)
    assert apply_no_code_floor(graded, rubric, {src.name: str(src)}) is graded


@pytest.mark.unit
def test_result_card_formats():
    from cqc_cpcc.result_card_export import (
        result_card_json,
        result_card_markdown,
        result_card_pdf,
        result_card_text,
    )

    result = _result(22)
    md = result_card_markdown("Ada Example", result, greeting_name="Ada Example")
    assert "# Result card: Ada Example" in md and "22 / 30 (73.3%)" in md
    assert "| Program Performance | 22/30 | Average | Works \\| mostly. |" in md
    assert "## Requirements" in md and "## Instructor notes" in md
    assert "**" not in result_card_text("Ada Example", result)
    assert json.loads(result_card_json("Ada Example", result))["result"]["total_points_earned"] == 22
    pdf = result_card_pdf("Ada Example", result)
    assert pdf.startswith(b"%PDF")
    import pymupdf
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        assert "Result card: Ada Example" in doc[0].get_text()


@pytest.mark.unit
def test_no_code_floor_reads_a_real_docx(tmp_path):
    import docx

    from cqc_cpcc.rubric_config import get_rubric_by_id
    from cqc_cpcc.rubric_grading import apply_no_code_floor

    report = tmp_path / "lab 4 super positon.docx"
    document = docx.Document()
    document.add_paragraph("Lab 4: superposition. Two waves add where they overlap.")
    document.save(report)
    rubric = get_rubric_by_id("csc134_cpp_exam_rubric")
    assert apply_no_code_floor(_result(4.5), rubric, {report.name: str(report)}).total_points_earned == 0

    with_code = tmp_path / "code.docx"
    document = docx.Document()
    for line in ("#include <iostream>", "int main() {", "  std::cout << 42;", "  return 0;", "}"):
        document.add_paragraph(line)
    document.save(with_code)
    graded = _result(18)
    assert apply_no_code_floor(graded, rubric, {with_code.name: str(with_code)}) is graded


@pytest.mark.unit
def test_student_drawer_is_the_dialog():
    import inspect

    from cqc_streamlit_app import grade_assignment
    source = inspect.getsource(grade_assignment)
    assert '@st.dialog("Student result", width="large")\ndef _student_drawer(' in source


@pytest.mark.unit
def test_no_closing_contains_a_number():
    import re
    for text in [t for _, t in CLOSING_BANDS] + [CLOSING_NO_WORK]:
        assert not re.search(r"\d", text)


@pytest.mark.unit
@pytest.mark.parametrize("prose", [
    "Lab report. In C++ we use std::cout to print.\nI wrote using namespace std at top.\nThe waves add up.",
    "My notes: Java prints with System.out.println and that is how output works.\nConclusion: it worked.",
])
def test_no_code_floor_is_not_fooled_by_prose_mentioning_code(tmp_path, prose):
    from cqc_cpcc.rubric_config import get_rubric_by_id
    from cqc_cpcc.rubric_grading import apply_no_code_floor

    report = tmp_path / "report.txt"
    report.write_text(prose)
    rubric = get_rubric_by_id("csc134_cpp_exam_rubric")
    assert apply_no_code_floor(_result(4.5), rubric, {report.name: str(report)}).total_points_earned == 0


@pytest.mark.unit
def test_inline_brightspace_feedback_ends_with_the_closing():
    from cqc_cpcc.utilities.brightspace_writeback import build_write_items_from_results
    items = build_write_items_from_results([("Ada Example", _result(29))])
    assert items[0].feedback_html.rstrip().endswith(f"<p>{dict(CLOSING_BANDS)[90]}</p>")


@pytest.mark.unit
def test_error_only_feedback_ends_with_a_matching_closing():
    from cqc_cpcc.exam_review import CodeGrader
    grader = CodeGrader(max_points=100, exam_instructions="x", exam_solution="x",
                        deduction_per_major_error=20, deduction_per_minor_error=5)
    grader.major_errors, grader.minor_errors = [], []
    assert grader.get_text_feedback().rstrip().endswith(dict(CLOSING_BANDS)[90])
    grader.invalid_reason = "No files were submitted."
    assert grader.get_text_feedback().rstrip().endswith(CLOSING_NO_WORK)


@pytest.mark.unit
def test_result_card_handles_missing_points_and_escapes_cells():
    from cqc_cpcc.result_card_export import result_card_markdown, result_card_pdf
    result = _result(10)
    result.criteria_results[0].points_earned = None
    result.criteria_results[0].criterion_name = "Name | with pipe\nand newline"
    md = result_card_markdown("Ada Example", result)
    assert "| Name \\| with pipe and newline | —/30 |" in md
    assert result_card_pdf("Ada Example", result).startswith(b"%PDF")


@pytest.mark.unit
def test_saved_runs_keep_greeting_names():
    from cqc_streamlit_app import results_store
    results_store.save_run("k1", "rubric", [("lab 4", _result(5))], [], greeting_names={"lab 4": ""})
    run = next(r for r in results_store.load_runs() if r["run_key"] == "k1")
    assert run["greeting_names"] == {"lab 4": ""}
